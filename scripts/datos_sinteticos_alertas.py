make run-apimake run-apimm#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Dataset sintético de alertas (FASE 0) — Phoenix Legal

Objetivo:
- Generar documentos con hechos verificables (fechas, importes, CIF/NIF, contrapartes, conceptos, IBAN, periodos)
  para disparar alertas en la pestaña de Alertas (sin diseñar a ciegas).

REGLAS DE ESTA FASE:
- No levanta UI ni API.
- No toca backend.
- No crea tests.
- No actualiza README/documentación del repo.

Catálogo inicial (para diseño posterior; NO implementa motor aquí):
- TGSS: deuda por periodos + apremio + falta de aplazamiento
- Banco: pagos a vinculadas + efectivo sin soporte + fraccionamientos
- Contabilidad: duplicados + mal pagada + descuadres/IVA
- Documentación: faltantes críticos (RNT/RLC, conciliaciones, etc.)
"""

from __future__ import annotations

import argparse
import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Tuple

from docx import Document as DocxDocument
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle


def _find_repo_root(start: Path) -> Path:
    """
    Encuentra root del repo para evitar rutas absolutas.
    Prioriza presencia de pyproject.toml o .git.
    """
    for p in [start, *start.parents]:
        if (p / "pyproject.toml").exists() or (p / ".git").exists():
            return p
    return start


REPO_ROOT = _find_repo_root(Path(__file__).resolve())

# Salida por defecto dentro del repo (portable).
DEFAULT_OUT_DIR = REPO_ROOT / "data" / "synthetic" / "alerts_dataset"

# Catálogo inicial (FASE 0: solo definición; no implementa motor aquí)
INITIAL_ALERT_CATALOG: dict[str, list[str]] = {
    "TGSS": ["deuda_por_periodos", "apremio", "falta_aplazamiento"],
    "BANCO": ["pagos_a_vinculadas", "efectivo_sin_soporte", "fraccionamientos"],
    "CONTABILIDAD": ["duplicados", "factura_mal_pagada", "descuadres_iva"],
    "DOCUMENTACION": ["faltantes_criticos_rnt_rlc", "faltantes_conciliaciones", "faltantes_balances"],
}

# Definition of Ready (dataset listo para diseñar sobre datos, no a ciegas)
DATASET_READY_CRITERIA: list[str] = [
    "PDFs principales con >=2 páginas (si ayuda a page_start/page_end).",
    "DOCX/EML/CSV con texto suficiente para chunks/snippets.",
    "Duplicados detectables por hash binario (no solo texto).",
    "Documentos navegables desde UI con trazabilidad (doc + ubicación).",
]


# =========================
# CONFIG (hechos duros)
# =========================

CASE_NAME = "RETAIL DEMO SL"
CIF = "B-12988731"
IBAN = "ES76 2100 4456 78 1234567890"

LINKED_COMPANY = "GRUPO XYZ SL"
LINKED_CIF = "B-99112233"

SUPPLIER_ALPHA = "PROVEEDOR ALPHA SL"
SUPPLIER_ALPHA_CIF = "B-55667788"

SUPPLIER_BETA = "PROVEEDOR BETA SL"
SUPPLIER_BETA_CIF = "B-11223344"

TGSS_PERIODS = [
    ("2023-01", 5200),
    ("2023-02", 5100),
    ("2023-03", 5300),
    ("2023-04", 5400),
    ("2023-05", 5200),
    ("2023-06", 5300),
    ("2023-07", 5400),
    ("2023-08", 5400),
]
TGSS_TOTAL = sum(v for _, v in TGSS_PERIODS)

BANK_MOVEMENTS_SEP_2023 = [
    ("2023-09-04", "Transferencia", "GRUPO XYZ SL", "Servicios", -6000.00, 23850.33),
    ("2023-09-08", "Transferencia", "PROVEEDOR ALPHA SL", "Factura FA-2023-077", -4500.00, 19350.33),
    ("2023-09-11", "Transferencia", "GRUPO XYZ SL", "Servicios", -6000.00, 13350.33),
    ("2023-09-14", "Retirada efectivo", "CAJERO", "Reintegro", -2000.00, 11350.33),
    ("2023-09-18", "Transferencia", "GRUPO XYZ SL", "Servicios", -6000.00, 5350.33),
    ("2023-09-21", "Transferencia", "NÓMINAS", "Nóminas septiembre", -4800.00, 550.33),
    ("2023-09-25", "Comisión", "BANCO", "Comisiones", -35.00, 515.33),
]

INVOICE_DUPLICATE = {
    "issuer": SUPPLIER_BETA,
    "issuer_cif": SUPPLIER_BETA_CIF,
    "invoice_no": "FB-2023-1021",
    "invoice_date": "2023-10-02",
    "concept": "Servicios de mantenimiento (septiembre 2023)",
    "net": 3200.00,
    "vat": 672.00,
    "total": 3872.00,
    "iban": "ES21 0081 1234 56 0000001234",
}

INVOICE_MISPAID = {
    "issuer": SUPPLIER_ALPHA,
    "issuer_cif": SUPPLIER_ALPHA_CIF,
    "invoice_no": "FA-2023-077",
    "invoice_date": "2023-09-05",
    "concept": "Suministro material tienda (agosto 2023)",
    "net": 4500.00,
    "vat": 945.00,
    "total": 5445.00,
    "paid_amount": 4500.00,  # pagada solo la base imponible (sin IVA)
    "pay_date": "2023-09-08",
}

LINKED_CONTRACT = {
    "provider": LINKED_COMPANY,
    "provider_cif": LINKED_CIF,
    "date": "2023-08-25",
    "concept": "Servicios de apoyo a gestión y consultoría",
    "monthly_fee": 6000.00,
    "deliverables": "No se detallan entregables concretos. Se indica 'soporte general'.",
}

EMAIL_BODY = (
    "Hola,\n\n"
    "He revisado los movimientos de septiembre y las transferencias a Grupo XYZ salen como 'servicios'.\n"
    "Para la asesoría laboral, de momento no enviéis el detalle de esos pagos hasta que cerremos el mes.\n"
    "Si necesitas soporte, lo vemos mañana.\n\n"
    "Gracias,\n"
    "Administración\n"
)


# =========================
# UTILS
# =========================


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def euros(x: float) -> str:
    s = f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return s + " €"


def _draw_header(c: canvas.Canvas, title: str, subtitle: str | None = None) -> None:
    c.setFont("Helvetica-Bold", 14)
    c.drawString(25 * mm, 287 * mm, title)
    c.setFont("Helvetica", 10)
    c.drawString(25 * mm, 281 * mm, CASE_NAME + f" · CIF {CIF}")
    if subtitle:
        c.drawString(25 * mm, 276 * mm, subtitle)
    c.setStrokeColor(colors.lightgrey)
    c.line(25 * mm, 273 * mm, 185 * mm, 273 * mm)


def _draw_footer(c: canvas.Canvas, page_num: int) -> None:
    c.setFont("Helvetica", 8)
    c.setFillColor(colors.grey)
    c.drawRightString(200 * mm, 10 * mm, f"Página {page_num}")
    c.setFillColor(colors.black)


def make_pdf(
    path: Path,
    *,
    title: str,
    subtitle: str,
    blocks: List[Tuple[str, List[str]]],
    table: List[List[str]] | None = None,
    min_pages: int = 1,
) -> None:
    """
    Genera PDF con cabecera/pie y opcionalmente fuerza un mínimo de páginas.
    Esto blinda la aparición de metadatos de página en evidencias cuando aplique.
    """
    c = canvas.Canvas(str(path), pagesize=A4)
    page = 1
    _draw_header(c, title, subtitle)

    y = 265 * mm
    c.setFont("Helvetica", 10)

    for (block_title, lines) in blocks:
        c.setFont("Helvetica-Bold", 10)
        c.drawString(25 * mm, y, block_title)
        y -= 6 * mm
        c.setFont("Helvetica", 10)

        for line in lines:
            if y < 40 * mm:
                _draw_footer(c, page)
                c.showPage()
                page += 1
                _draw_header(c, title, subtitle)
                y = 265 * mm
                c.setFont("Helvetica", 10)
            c.drawString(25 * mm, y, line)
            y -= 5.5 * mm

        y -= 4 * mm

    if table:
        if y < 90 * mm:
            _draw_footer(c, page)
            c.showPage()
            page += 1
            _draw_header(c, title, subtitle)
            y = 265 * mm

        t = Table(table, colWidths=[40 * mm, 40 * mm, 40 * mm, 50 * mm])
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0f0")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                    ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                    ("ALIGN", (0, 0), (-1, 0), "LEFT"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )

        tw, th = t.wrapOn(c, 160 * mm, 120 * mm)
        t.drawOn(c, 25 * mm, y - th)
        y = y - th - 10 * mm

    # Forzar páginas mínimas si hace falta
    while page < max(1, int(min_pages)):
        _draw_footer(c, page)
        c.showPage()
        page += 1
        _draw_header(c, title, subtitle)
        c.setFont("Helvetica", 10)
        c.setFillColor(colors.black)
        c.drawString(25 * mm, 260 * mm, "ANEXO — Página adicional para pruebas de metadatos (page_start/page_end).")
        c.setFont("Helvetica", 9)
        c.setFillColor(colors.grey)
        c.drawString(25 * mm, 252 * mm, "Contenido sintético. No concluye intención. Solo soporta trazabilidad.")
        c.setFillColor(colors.black)

    _draw_footer(c, page)
    c.save()


def make_eml(
    path: Path,
    *,
    subject: str,
    from_name: str,
    from_email: str,
    to_email: str,
    sent: str,
    body: str,
) -> None:
    content = (
        f"From: {from_name} <{from_email}>\n"
        f"To: {to_email}\n"
        f"Subject: {subject}\n"
        f"Date: {sent}\n"
        f"MIME-Version: 1.0\n"
        f"Content-Type: text/plain; charset=utf-8\n"
        f"\n"
        f"{body}\n"
    )
    path.write_text(content, encoding="utf-8")


def make_docx_letter(path: Path, title: str, paragraphs: List[str]) -> None:
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(path))


def ensure_dirs(out_dir: Path) -> Dict[str, Path]:
    paths = {
        "tgss": out_dir / "01_tgss",
        "banco": out_dir / "02_banco",
        "contabilidad": out_dir / "03_contabilidad",
        "facturas": out_dir / "03_contabilidad" / "facturas",
        "vinculadas": out_dir / "04_operaciones_vinculadas",
        "emails": out_dir / "05_emails",
        "nominas": out_dir / "06_nominas_seguros_sociales",
        "manifest": out_dir,
    }
    for p in paths.values():
        p.mkdir(parents=True, exist_ok=True)
    return paths


# =========================
# GENERATORS
# =========================


def generate_tgss(paths: Dict[str, Path]) -> Path:
    pdf_path = paths["tgss"] / "TGSS_Providencia_Apremio_Exp123.pdf"
    table = [["Periodo", "Importe", "Recargo", "Total periodo"]]
    for period, amount in TGSS_PERIODS:
        recargo = round(amount * 0.10, 2)
        total = round(amount + recargo, 2)
        table.append([period, euros(amount), euros(recargo), euros(total)])

    blocks = [
        (
            "Datos del expediente",
            [
                "Organismo: Tesorería General de la Seguridad Social (TGSS)",
                "Nº expediente: 123/2023/APR",
                "Fecha de emisión: 05/09/2023",
                "Fecha de notificación: 10/09/2023",
                f"Deudor: {CASE_NAME} · CIF {CIF}",
                "Estado: Providencia de apremio (cuotas no ingresadas)",
            ],
        ),
        (
            "Resumen",
            [
                f"Importe principal reclamado: {euros(TGSS_TOTAL)}",
                "Periodo reclamado: 2023-01 a 2023-08 (8 periodos consecutivos).",
                "Observación: No consta ingreso de cuotas en los periodos indicados.",
            ],
        ),
        (
            "Notas",
            [
                "Documento sintético para pruebas de ingestión y alertas.",
                "Incluye fechas, periodos e importes verificables.",
            ],
        ),
    ]
    make_pdf(
        pdf_path,
        title="Providencia de Apremio (TGSS)",
        subtitle="Expediente 123/2023/APR",
        blocks=blocks,
        table=table,
        min_pages=2,  # blindaje: asegurar páginas para metadatos
    )
    return pdf_path


def generate_bank_extract(paths: Dict[str, Path]) -> Path:
    pdf_path = paths["banco"] / "Extracto_Bancario_Sep2023.pdf"
    table = [["Fecha", "Tipo", "Contraparte", "Importe / Saldo"]]
    for (d, typ, counterparty, concept, amount, saldo) in BANK_MOVEMENTS_SEP_2023:
        imp = euros(abs(amount))
        sign = "-" if amount < 0 else "+"
        table.append([d, f"{typ} · {concept}", counterparty, f"{sign}{imp} · saldo {euros(saldo)}"])

    blocks = [
        (
            "Datos de cuenta",
            [
                f"Titular: {CASE_NAME}",
                f"CIF: {CIF}",
                f"IBAN: {IBAN}",
                "Entidad: Banco Sintético S.A.",
                "Periodo: 01/09/2023 a 30/09/2023",
            ],
        ),
        (
            "Observaciones",
            [
                f"Se registran transferencias recurrentes a {LINKED_COMPANY} ({LINKED_CIF}) con concepto 'Servicios'.",
                "Se registra al menos una retirada de efectivo sin soporte documental asociado en este extracto.",
            ],
        ),
    ]
    make_pdf(
        pdf_path,
        title="Extracto Bancario",
        subtitle="Movimientos septiembre 2023",
        blocks=blocks,
        table=table,
        min_pages=2,  # blindaje: asegurar páginas para metadatos
    )
    return pdf_path


def generate_invoice_pdf(path: Path, inv: Dict, filename_title: str) -> Path:
    blocks = [
        (
            "Emisor",
            [
                f"Proveedor: {inv['issuer']}",
                f"CIF: {inv['issuer_cif']}",
                "Dirección: Calle Industria 12, 28000 Madrid",
            ],
        ),
        (
            "Receptor",
            [
                f"Cliente: {CASE_NAME}",
                f"CIF: {CIF}",
                "Dirección: Calle Mayor 1, 28000 Madrid",
            ],
        ),
        (
            "Factura",
            [
                f"Nº factura: {inv['invoice_no']}",
                f"Fecha factura: {inv['invoice_date']}",
                f"Concepto: {inv['concept']}",
            ],
        ),
        (
            "Importes",
            [
                f"Base imponible: {euros(inv['net'])}",
                f"IVA (21%): {euros(inv['vat'])}",
                f"TOTAL: {euros(inv['total'])}",
                f"IBAN proveedor: {inv.get('iban', 'N/D')}",
            ],
        ),
        (
            "Notas",
            [
                "Documento sintético para pruebas de alertas de duplicados y contabilidad.",
            ],
        ),
    ]
    make_pdf(
        path,
        title=f"Factura {filename_title}",
        subtitle=f"{inv['invoice_no']}",
        blocks=blocks,
        table=None,
        min_pages=1,
    )
    return path


def generate_mispaid_support(paths: Dict[str, Path]) -> Path:
    pdf_path = paths["contabilidad"] / "Justificante_Pago_FA-2023-077.pdf"
    blocks = [
        (
            "Resumen del pago",
            [
                f"Factura: {INVOICE_MISPAID['invoice_no']} ({INVOICE_MISPAID['invoice_date']})",
                f"Proveedor: {INVOICE_MISPAID['issuer']} · CIF {INVOICE_MISPAID['issuer_cif']}",
                f"Total factura: {euros(INVOICE_MISPAID['total'])}",
                f"Importe pagado: {euros(INVOICE_MISPAID['paid_amount'])}",
                f"Fecha de pago: {INVOICE_MISPAID['pay_date']}",
                "Observación: Se ha abonado únicamente la base imponible (sin IVA).",
            ],
        ),
        (
            "Detalle",
            [
                "Posible causa: error administrativo en la orden de pago o discrepancia en validación de IVA.",
                "Acción recomendada: conciliar con proveedor, solicitar rectificativa o completar pago si procede.",
            ],
        ),
    ]
    make_pdf(
        pdf_path,
        title="Justificante de Pago",
        subtitle="Discrepancia importe vs factura",
        blocks=blocks,
        table=None,
        min_pages=1,
    )
    return pdf_path


def generate_linked_party_contract(paths: Dict[str, Path]) -> Path:
    docx_path = paths["vinculadas"] / "Contrato_Servicios_GrupoXYZ_2023-08-25.docx"
    paragraphs = [
        f"Fecha: {LINKED_CONTRACT['date']}",
        f"Entre: {CASE_NAME} (CIF {CIF}) y {LINKED_CONTRACT['provider']} (CIF {LINKED_CONTRACT['provider_cif']}).",
        f"Objeto: {LINKED_CONTRACT['concept']}.",
        f"Retribución: cuota mensual de {euros(LINKED_CONTRACT['monthly_fee'])}.",
        f"Entregables: {LINKED_CONTRACT['deliverables']}",
        "Duración: 6 meses renovables.",
        "Observación: Contrato sintético para pruebas; texto deliberadamente poco específico.",
    ]
    make_docx_letter(docx_path, "Contrato de Servicios", paragraphs)
    return docx_path


def generate_email(paths: Dict[str, Path]) -> Path:
    eml_path = paths["emails"] / "Email_Admin_Asesoria_2023-09-22.eml"
    make_eml(
        eml_path,
        subject="Cierre septiembre - movimientos y asesoría",
        from_name="Administración",
        from_email="admin@retaildemosl.local",
        to_email="gerencia@retaildemosl.local",
        sent="Fri, 22 Sep 2023 10:15:00 +0200",
        body=EMAIL_BODY,
    )
    return eml_path


def generate_nominas_placeholder(paths: Dict[str, Path]) -> Tuple[Path, Path]:
    pdf_nomina = paths["nominas"] / "Nominas_Sep2023_Resumen.pdf"
    blocks = [
        (
            "Resumen nóminas",
            [
                "Periodo: septiembre 2023",
                "Nº trabajadores: 6",
                "Total netos abonados: 4.800,00 €",
                "Forma de pago: transferencia (según extracto bancario).",
            ],
        ),
        (
            "Nota",
            [
                "Documento sintético. No incluye datos personales reales.",
                "Se usa para cruzar con alertas TGSS (cuotas no ingresadas vs nóminas abonadas).",
            ],
        ),
    ]
    make_pdf(
        pdf_nomina,
        title="Resumen de Nóminas",
        subtitle="Septiembre 2023",
        blocks=blocks,
        table=None,
        min_pages=1,
    )

    txt_missing = paths["nominas"] / "FALTA_RNT_RLC_2023-01_a_2023-08.txt"
    txt_missing.write_text(
        "ALERTA DOCUMENTAL (sintética)\n\n"
        "No consta en el expediente la Relación Nominal de Trabajadores (RNT) / Relación de Liquidación de Cotizaciones (RLC)\n"
        "para los periodos 2023-01 a 2023-08.\n",
        encoding="utf-8",
    )
    return pdf_nomina, txt_missing


def generate_invoices(paths: Dict[str, Path]) -> List[Path]:
    out: List[Path] = []

    inv1 = paths["facturas"] / "Factura_FB-2023-1021_PROV_BETA.pdf"
    inv2 = paths["facturas"] / "Factura_Mantenimiento_Sep2023_Beta.pdf"  # distinto nombre, mismo binario

    generate_invoice_pdf(inv1, INVOICE_DUPLICATE, "PROVEEDOR BETA")
    # Blindaje duplicado: copia exacta de bytes (hash binario idéntico)
    inv2.write_bytes(inv1.read_bytes())

    out.extend([inv1, inv2])

    inv_m = paths["facturas"] / "Factura_FA-2023-077_PROV_ALPHA.pdf"
    generate_invoice_pdf(
        inv_m,
        {
            "issuer": INVOICE_MISPAID["issuer"],
            "issuer_cif": INVOICE_MISPAID["issuer_cif"],
            "invoice_no": INVOICE_MISPAID["invoice_no"],
            "invoice_date": INVOICE_MISPAID["invoice_date"],
            "concept": INVOICE_MISPAID["concept"],
            "net": INVOICE_MISPAID["net"],
            "vat": INVOICE_MISPAID["vat"],
            "total": INVOICE_MISPAID["total"],
            "iban": "ES11 2100 0000 00 9999999999",
        },
        "PROVEEDOR ALPHA",
    )
    out.append(inv_m)
    out.append(generate_mispaid_support(paths))
    return out


def generate_accounting_csv(paths: Dict[str, Path]) -> Path:
    csv_path = paths["contabilidad"] / "Libro_Mayor_2023_Q3.csv"
    rows = [
        "fecha,asiento,cuenta,descripcion,tercero,tercero_cif,debe,haber",
        "2023-09-08,3201,600000,Compra material tienda,PROVEEDOR ALPHA SL,B-55667788,4500.00,0.00",
        # IVA soportado mal recogido (para provocar descuadre)
        "2023-09-08,3202,472000,IVA soportado,PROVEEDOR ALPHA SL,B-55667788,0.00,945.00",
        # Pagos vinculadas con concepto genérico
        "2023-09-04,3190,623000,Servicios,GRUPO XYZ SL,B-99112233,6000.00,0.00",
        "2023-09-11,3210,623000,Servicios,GRUPO XYZ SL,B-99112233,6000.00,0.00",
        "2023-09-18,3220,623000,Servicios,GRUPO XYZ SL,B-99112233,6000.00,0.00",
        # Retirada efectivo sin soporte
        "2023-09-14,3208,570000,Reintegro cajero,CAJERO,,2000.00,0.00",
    ]
    csv_path.write_text("\n".join(rows) + "\n", encoding="utf-8")
    return csv_path


def _pdf_page_count(path: Path) -> int | None:
    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(str(path))
        return len(reader.pages)
    except Exception:
        return None


def write_manifest(out_dir: Path, created_files: List[Path]) -> Path:
    """
    Manifest interno del dataset (no es README del repo).
    Incluye hashes y checks mínimos (Definition of Ready).
    """
    manifest_path = out_dir / "MANIFEST_ALERTAS_DATASET.txt"
    lines: List[str] = []
    lines.append("DATASET SINTÉTICO ALERTAS — MANIFEST")
    lines.append(f"OUT_DIR: {out_dir}")
    lines.append("")
    lines.append("ARCHIVOS (ruta relativa + sha256):")
    hashes: Dict[str, List[Path]] = {}
    for p in created_files:
        if not p.is_file():
            continue
        h = sha256_file(p)
        hashes.setdefault(h, []).append(p)
        rel = p.relative_to(out_dir)
        lines.append(f"- {rel}  (sha256: {h})")

    lines.append("")
    dups = {h: ps for h, ps in hashes.items() if len(ps) > 1}
    if dups:
        lines.append("DUPLICADOS (mismo sha256):")
        for h, ps in dups.items():
            lines.append(f"- {h[:12]}…")
            for p in ps:
                lines.append(f"  - {p.relative_to(out_dir)}")
    else:
        lines.append("DUPLICADOS: No se detectaron duplicados por hash binario.")

    lines.append("")
    lines.append("CHECKS (Definition of Ready):")
    # PDFs principales
    tgss_pdf = out_dir / "01_tgss" / "TGSS_Providencia_Apremio_Exp123.pdf"
    bank_pdf = out_dir / "02_banco" / "Extracto_Bancario_Sep2023.pdf"
    for label, pdf in [("TGSS", tgss_pdf), ("BANCO", bank_pdf)]:
        if pdf.exists():
            pc = _pdf_page_count(pdf)
            if pc is None:
                lines.append(f"- PDF {label}: no se pudo contar páginas (PyPDF2 no disponible).")
            else:
                lines.append(f"- PDF {label}: {pc} página(s).")
        else:
            lines.append(f"- PDF {label}: NO ENCONTRADO.")

    # Texto suficiente en EML/CSV/DOCX
    eml = out_dir / "05_emails" / "Email_Admin_Asesoria_2023-09-22.eml"
    csv = out_dir / "03_contabilidad" / "Libro_Mayor_2023_Q3.csv"
    docx = out_dir / "04_operaciones_vinculadas" / "Contrato_Servicios_GrupoXYZ_2023-08-25.docx"
    for label, fp in [("EML", eml), ("CSV", csv), ("DOCX", docx)]:
        if fp.exists():
            sz = fp.stat().st_size
            lines.append(f"- {label}: OK (tamaño {sz} bytes).")
        else:
            lines.append(f"- {label}: NO ENCONTRADO.")

    manifest_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return manifest_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera dataset sintético de alertas (FASE 0).")
    parser.add_argument(
        "--out-dir",
        default=str(DEFAULT_OUT_DIR),
        help="Directorio de salida dentro del repo (default: data/synthetic/alerts_dataset).",
    )
    args = parser.parse_args()

    out_dir = Path(args.out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    paths = ensure_dirs(out_dir)

    created: List[Path] = []
    created.append(generate_tgss(paths))
    created.append(generate_bank_extract(paths))
    created.extend(generate_invoices(paths))
    created.append(generate_linked_party_contract(paths))
    created.append(generate_email(paths))
    created.extend(list(generate_nominas_placeholder(paths)))
    created.append(generate_accounting_csv(paths))
    created.append(write_manifest(out_dir, created))

    print("✅ Dataset de alertas generado")
    print(f"📁 Ruta: {out_dir}")
    print(f"📄 Total archivos: {len(created)}")

    # Prueba de duplicados por hash binario
    hashes: Dict[str, List[Path]] = {}
    for p in created:
        if p.is_file():
            h = sha256_file(p)
            hashes.setdefault(h, []).append(p)
    dups = {h: ps for h, ps in hashes.items() if len(ps) > 1}
    if dups:
        print("🧬 Duplicados (mismo sha256):")
        for h, ps in dups.items():
            print(" -", h[:12], "=>")
            for p in ps:
                print("   ", p.relative_to(out_dir))
    else:
        print("🧬 No se detectaron duplicados por hash binario.")

    # Pages check (opcional, pero útil)
    tgss_pdf = out_dir / "01_tgss" / "TGSS_Providencia_Apremio_Exp123.pdf"
    bank_pdf = out_dir / "02_banco" / "Extracto_Bancario_Sep2023.pdf"
    for label, pdf in [("TGSS", tgss_pdf), ("BANCO", bank_pdf)]:
        if pdf.exists():
            pc = _pdf_page_count(pdf)
            if pc is not None:
                print(f"📄 PDF {label}: {pc} página(s)")


if __name__ == "__main__":
    main()

