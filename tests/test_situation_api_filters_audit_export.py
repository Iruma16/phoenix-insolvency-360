import io
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.database import Base, get_db
from app.main import app
from app.models.case import Case
from app.models.document import Document


@pytest.fixture
def client_with_db():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    session = Session()

    case = Case(case_id="case_s2", name="Caso S2")
    session.add(case)
    session.commit()

    doc = Document(
        document_id="doc_s2",
        case_id=case.case_id,
        filename="aeat.pdf",
        sha256_hash="0" * 64,
        file_size_bytes=1,
        mime_type="application/pdf",
        uploaded_at=datetime.utcnow(),
        doc_type="AEAT",
        source="test",
        date_start=datetime.utcnow(),
        date_end=datetime.utcnow(),
        reliability="original",
        file_format="pdf",
        storage_path="/tmp/aeat.pdf",
        created_at=datetime.utcnow(),
    )
    session.add(doc)
    session.commit()

    def _override_get_db():
        try:
            yield session
        finally:
            pass

    app.dependency_overrides[get_db] = _override_get_db
    try:
        c = TestClient(app)
        # Exponer sesión para tests que necesitan seed adicional sin endpoints.
        setattr(c, "_db_session", session)
        yield c
    finally:
        app.dependency_overrides.pop(get_db, None)
        session.close()


def _evidence():
    return [{"document_id": "doc_s2", "page": 1, "note": "Consta en documento aportado."}]


def test_rejects_evidence_document_from_other_case(client_with_db):
    # Crear otro caso y un documento en ese otro caso (seed directo en DB)
    db = getattr(client_with_db, "_db_session")
    other = Case(case_id="case_other", name="Otro caso")
    db.add(other)
    db.commit()

    other_doc = Document(
        document_id="doc_other",
        case_id=other.case_id,
        filename="otro.pdf",
        sha256_hash="1" * 64,
        file_size_bytes=1,
        mime_type="application/pdf",
        uploaded_at=datetime.utcnow(),
        doc_type="OTRO",
        source="test",
        date_start=datetime.utcnow(),
        date_end=datetime.utcnow(),
        reliability="original",
        file_format="pdf",
        storage_path="/tmp/otro.pdf",
        created_at=datetime.utcnow(),
    )
    db.add(other_doc)
    db.commit()

    # Intentar usar document_id de OTRO caso como evidencia en case_s2 => debe rechazar
    bad = client_with_db.post(
        "/api/cases/case_s2/situation/invoices",
        json={
            "created_by": "abogado",
            "reason": "Alta con evidencia incorrecta",
            "evidence": [{"document_id": "doc_other", "page": 1, "note": "No pertenece al caso."}],
            "supplier": "Proveedor SL",
            "amount_total": 100.0,
            "status": "pendiente",
        },
    )
    assert bad.status_code == 400, bad.text


def test_list_invoices_paginado_total(client_with_db):
    # Crear 3 facturas
    for i in range(3):
        r = client_with_db.post(
            "/api/cases/case_s2/situation/invoices",
            json={
                "created_by": "abogado",
                "reason": f"Alta inicial {i} con evidencia suficiente",
                "evidence": _evidence(),
                "supplier": f"Proveedor {i}",
                "amount_total": 100.0 + i,
                "status": "pendiente",
            },
        )
        assert r.status_code == 200

    # Página 1, 2 items
    r = client_with_db.get(
        "/api/cases/case_s2/situation/invoices",
        params={"page": 1, "page_size": 2},
    )
    assert r.status_code == 200
    data = r.json()
    assert data["total"] == 3
    assert data["page"] == 1
    assert data["page_size"] == 2
    assert len(data["items"]) == 2


def test_evidence_and_audit_endpoints(client_with_db):
    # Crear factura
    r = client_with_db.post(
        "/api/cases/case_s2/situation/invoices",
        json={
            "created_by": "abogado",
            "reason": "Alta inicial",
            "evidence": _evidence(),
            "supplier": "Proveedor SL",
            "amount_total": 100.0,
        },
    )
    assert r.status_code == 200
    inv = r.json()
    logical_id = inv["logical_id"]
    record_id = inv["record_id"]

    # Ver evidencia del record
    ev = client_with_db.get(f"/api/cases/case_s2/situation/INVOICE/{record_id}/evidence")
    assert ev.status_code == 200
    ev_items = ev.json()["items"]
    assert len(ev_items) == 1
    assert ev_items[0]["document_id"] == "doc_s2"

    # Añadir evidencia extra
    add = client_with_db.post(
        f"/api/cases/case_s2/situation/INVOICE/{record_id}/evidence",
        json={
            "created_by": "abogado",
            "reason": "Se añade soporte adicional",
            "evidence": [
                {"document_id": "doc_s2", "page": 2, "note": "Segundo soporte probatorio."}
            ],
        },
    )
    assert add.status_code == 200
    assert len(add.json()["items"]) == 2

    # Ver auditoría del logical_id (debe incluir ADD_EVIDENCE)
    au = client_with_db.get(f"/api/cases/case_s2/situation/INVOICE/{logical_id}/audit")
    assert au.status_code == 200
    actions = [x["action"] for x in au.json()["items"]]
    assert "CREATE" in actions
    assert "ADD_EVIDENCE" in actions


def test_export_multi_sheet_excel(client_with_db):
    # Crear una factura para que no exporte vacío
    r = client_with_db.post(
        "/api/cases/case_s2/situation/invoices",
        json={
            "created_by": "abogado",
            "reason": "Alta para export",
            "evidence": _evidence(),
            "supplier": "Proveedor SL",
            "amount_total": 100.0,
        },
    )
    assert r.status_code == 200

    resp = client_with_db.get("/api/cases/case_s2/situation/export.xlsx")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers.get(
        "content-type", ""
    )

    wb = load_workbook(filename=io.BytesIO(resp.content))
    assert "Facturas" in wb.sheetnames
    assert "Créditos" in wb.sheetnames
    assert "Bienes" in wb.sheetnames
    assert "Deuda pública" in wb.sheetnames
    assert "Juzgado" in wb.sheetnames
